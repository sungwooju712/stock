# -*- coding: utf-8 -*-
"""Excel report generation in the style of the supplied KOSPI workbook.

Five worksheets are produced for a watchlist report:
    한장보고 / 전체보고 / 일일종가 / 월별종가 / 연간데이터
Native Excel line/bar charts are embedded, and the data remains editable.
"""

from __future__ import annotations

import math
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.worksheet.hyperlink import Hyperlink

from .constants import CHART_COLORS, COLORS, DISCLAIMER
from .formatters import fmt_market_cap, fmt_pct, fmt_price
from .paths import naver_stock_url
from .services import MarketService, ScreenResult


# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------
FONT_NAME = "맑은 고딕"
NAVY = "0B1F33"
NAVY_2 = "14395A"
TEAL = "0F9D8A"
TEAL_DARK = "0A7C6D"
AMBER = "F2A93B"
PALE_BLUE = "EAF2F8"
PALE_TEAL = "E7F8F4"
PALE_AMBER = "FFF5E4"
ROW_ALT = "F7FAFC"
BORDER_COLOR = "D7E1EA"
TEXT = "122033"
MUTED = "64748B"
RISE = "D93636"
FALL = "2563EB"


def font(size: int = 10, bold: bool = False, color: str = TEXT, italic: bool = False) -> Font:
    return Font(name=FONT_NAME, size=size, bold=bold, color=color, italic=italic)


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


THIN = Side(style="thin", color=BORDER_COLOR)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
WRAP_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

PRICE_FMT = "#,##0"
PERCENT_SIGNED_FMT = '+0.00"%";-0.00"%"'
PERCENT_FMT = '0.00"%"'
RATIO_FMT = '0.00"배"'
EOK_FMT = '#,##0"억원"'
COUNT_FMT = "#,##0"


class ExcelReport:
    def __init__(self, service: MarketService):
        self.service = service

    # ------------------------------------------------------------------
    # Public exports
    # ------------------------------------------------------------------
    def export_watchlist(self, path: str | Path, codes: Sequence[str] | None = None) -> Path:
        bundle = self.service.report_bundle(codes)
        stocks = bundle.get("stocks") or []
        if not stocks:
            raise ValueError("내보낼 관심종목이 없습니다.")
        workbook = Workbook()
        workbook.remove(workbook.active)
        workbook.properties.title = "KOSPI 종목 종가 분석 보고서"
        workbook.properties.subject = "2023년~전일종가 종가 분석"
        workbook.properties.creator = "KOSPI 인사이트"
        workbook.properties.keywords = "KOSPI, 종가, PER, 부채비율, 영업이익률, 스크리닝"

        datasets = self._prepare_datasets(bundle)
        # Data sheets must exist before the chart-bearing summary sheets are built.
        self._build_daily_sheet(workbook, bundle, datasets)
        self._build_monthly_sheet(workbook, bundle, datasets)
        self._build_annual_sheet(workbook, bundle, datasets)
        self._build_summary_sheet(workbook, bundle, datasets)
        self._build_detail_sheet(workbook, bundle, datasets)
        workbook.move_sheet("한장보고", offset=-3)
        workbook.move_sheet("전체보고", offset=-3)
        workbook.calculation.fullCalcOnLoad = True
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        return path

    def export_screen(self, path: str | Path, result: ScreenResult) -> Path:
        if not result.rows:
            raise ValueError("내보낼 스크리닝 결과가 없습니다.")
        workbook = Workbook()
        ws = workbook.active
        ws.title = "스크리닝 결과"
        ws.sheet_view.showGridLines = False
        max_col = 17
        self._title_block(
            ws,
            "KOSPI 스크리닝 결과",
            self._screen_subtitle(result),
            "출처: NAVER Finance / 한국거래소(KRX) | 재무 기준: 최근 확정 연차 | 투자 권유 아님",
            max_col,
        )
        criteria_text = self._criteria_text(result)
        ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=max_col)
        cell = ws.cell(5, 1, f"적용 조건  |  {criteria_text}")
        cell.font = font(9.5, True, NAVY_2)
        cell.fill = fill(PALE_TEAL)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.row_dimensions[5].height = 34
        ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=max_col)
        summary = ws.cell(
            6,
            1,
            f"통과 {len(result.rows):,}개 / KOSPI 전체 {result.total:,}개  ·  "
            f"재무데이터 보유 {result.fundamental_ready:,}개  ·  "
            f"평균 PER {fmt_pct(result.averages.get('per'), digits=2) if result.averages.get('per') is not None else '–'}",
        )
        summary.font = font(9, False, MUTED)
        summary.alignment = LEFT
        summary.fill = fill(ROW_ALT)
        summary.border = BORDER
        ws.row_dimensions[6].height = 22

        headers = [
            "순위",
            "종목명",
            "코드",
            "현재가(원)",
            "등락률",
            "시가총액(억원)",
            "PER(배)",
            "PBR(배)",
            "ROE",
            "영업이익률",
            "부채비율",
            "배당수익률",
            "매출액(억원)",
            "당기순이익(억원)",
            "52주 최고(원)",
            "52주 최저(원)",
            "관심종목",
        ]
        header_row = 8
        for column, label in enumerate(headers, start=1):
            cell = ws.cell(header_row, column, label)
            cell.font = font(9, True, "FFFFFF")
            cell.fill = fill(NAVY_2)
            cell.alignment = CENTER
            cell.border = BORDER
        ws.row_dimensions[header_row].height = 30

        for index, row in enumerate(result.rows, start=1):
            excel_row = header_row + index
            values = [
                row.get("rank"),
                row.get("name"),
                row.get("code"),
                row.get("close"),
                row.get("change_pct"),
                round((row.get("market_cap") or 0) / 100_000_000, 1),
                row.get("per"),
                row.get("pbr"),
                row.get("roe"),
                row.get("op_margin"),
                row.get("debt_ratio"),
                row.get("dividend_yield"),
                row.get("sales"),
                row.get("net_income"),
                row.get("high52"),
                row.get("low52"),
                "★" if row.get("in_watchlist") else "",
            ]
            row_fill = fill("FFFFFF" if index % 2 else ROW_ALT)
            for column, value in enumerate(values, start=1):
                cell = ws.cell(excel_row, column, value)
                cell.font = font(9, False)
                cell.fill = row_fill
                cell.border = BORDER
                cell.alignment = CENTER if column not in (2, 3) else LEFT
                if column == 5:
                    number = value if isinstance(value, (int, float)) else 0
                    cell.font = font(9, True, RISE if number > 0 else FALL if number < 0 else MUTED)
                    cell.number_format = PERCENT_SIGNED_FMT
                elif column in (4, 6, 13, 14, 15, 16):
                    cell.number_format = COUNT_FMT
                    cell.alignment = RIGHT
                elif column == 7:
                    cell.number_format = RATIO_FMT
                elif column == 8:
                    cell.number_format = RATIO_FMT
                elif column in (9, 10, 11, 12):
                    cell.number_format = PERCENT_FMT
                elif column == 3:
                    cell.hyperlink = Hyperlink(ref=cell.coordinate, target=naver_stock_url(row["code"]))
                    cell.font = font(9, False, FALL)
                elif column == 2:
                    cell.alignment = LEFT
                elif column == 17:
                    cell.font = font(9, True, AMBER)
            ws.row_dimensions[excel_row].height = 20
        last_row = header_row + len(result.rows)
        ws.auto_filter.ref = f"A{header_row}:Q{last_row}"
        ws.freeze_panes = f"A{header_row + 1}"
        widths = [6, 18, 10, 13, 10, 17, 10, 10, 10, 12, 12, 12, 15, 16, 14, 14, 9]
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
        self._page_setup(ws)

        note_row = last_row + 2
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=max_col)
        note = ws.cell(note_row, 1, DISCLAIMER)
        note.font = font(8, False, MUTED)
        note.alignment = WRAP_LEFT
        note.fill = fill(ROW_ALT)
        workbook.calculation.fullCalcOnLoad = True
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        return path

    # ------------------------------------------------------------------
    # Data preparation
    # ------------------------------------------------------------------
    def _prepare_datasets(self, bundle: dict[str, Any]) -> dict[str, Any]:
        stocks = bundle["stocks"]
        codes = [stock["code"] for stock in stocks]
        raw_series: dict[str, list[dict[str, Any]]] = bundle["series"]
        date_set = {row["date"] for rows in raw_series.values() for row in rows}
        dates = sorted(date_set)
        lookup = {
            code: {row["date"]: row["close"] for row in raw_series.get(code, [])}
            for code in codes
        }
        daily_dates = dates
        months = sorted({day[:7] for day in dates})
        month_lookup: dict[str, dict[str, int]] = {}
        for code in codes:
            per_month: dict[str, int] = {}
            for row in raw_series.get(code, []):
                per_month[row["date"][:7]] = row["close"]
            month_lookup[code] = per_month
        years = sorted({int(day[:4]) for day in dates})
        annual_lookup: dict[str, dict[int, int]] = {}
        for code in codes:
            per_year: dict[int, int] = {}
            for row in raw_series.get(code, []):
                per_year[int(row["date"][:4])] = row["close"]
            annual_lookup[code] = per_year
        annual_returns: dict[str, dict[int, float | None]] = {}
        for code in codes:
            values = annual_lookup[code]
            result: dict[int, float | None] = {}
            for year in years:
                previous = values.get(year - 1)
                current = values.get(year)
                result[year] = (current / previous - 1) * 100 if current and previous else None
            annual_returns[code] = result
        return {
            "dates": daily_dates,
            "months": months,
            "years": years,
            "lookup": lookup,
            "month_lookup": month_lookup,
            "annual_lookup": annual_lookup,
            "annual_returns": annual_returns,
        }

    # ------------------------------------------------------------------
    # Workbook styling helpers
    # ------------------------------------------------------------------
    def _title_block(self, ws, title: str, subtitle: str, source: str, max_col: int) -> None:
        ws.sheet_view.showGridLines = False
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_col)
        cell = ws.cell(1, 1, title)
        cell.font = font(17, True, "FFFFFF")
        cell.fill = fill(NAVY)
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 36
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_col)
        cell = ws.cell(2, 1, subtitle)
        cell.font = font(9.5, True, NAVY_2)
        cell.fill = fill(PALE_BLUE)
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[2].height = 22
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=max_col)
        cell = ws.cell(3, 1, source)
        cell.font = font(8, False, MUTED)
        cell.fill = fill("FFFFFF")
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[3].height = 19
        self._page_setup(ws)

    def _page_setup(self, ws) -> None:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = 0.35
        ws.page_margins.right = 0.35
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5

    def _section_header(self, ws, row: int, text: str, max_col: int) -> None:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_col)
        cell = ws.cell(row, 1, text)
        cell.font = font(11, True, NAVY_2)
        cell.fill = fill(PALE_TEAL)
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        cell.border = BORDER
        ws.row_dimensions[row].height = 24

    def _table_header(self, ws, row: int, labels: Sequence[str], *, start_col: int = 1) -> None:
        for index, label in enumerate(labels, start=start_col):
            cell = ws.cell(row, index, label)
            cell.font = font(9, True, "FFFFFF")
            cell.fill = fill(NAVY_2)
            cell.alignment = CENTER
            cell.border = BORDER
        ws.row_dimensions[row].height = 24

    def _label_cell(self, ws, row: int, text: str, *, fill_color: str = "FFFFFF") -> None:
        cell = ws.cell(row, 1, text)
        cell.font = font(9, True, TEXT)
        cell.fill = fill(fill_color)
        cell.alignment = LEFT
        cell.border = BORDER

    def _value_cell(
        self,
        ws,
        row: int,
        column: int,
        value: Any,
        *,
        number_format: str | None = None,
        alignment: Alignment = RIGHT,
        bold: bool = False,
        color: str | None = None,
        fill_color: str = "FFFFFF",
    ):
        cell = ws.cell(row, column, value)
        cell.font = font(9, bold, color or TEXT)
        cell.fill = fill(fill_color)
        cell.alignment = alignment
        cell.border = BORDER
        if number_format:
            cell.number_format = number_format
        return cell

    # ------------------------------------------------------------------
    # Sheet 1: 한장보고
    # ------------------------------------------------------------------
    def _build_summary_sheet(self, wb: Workbook, bundle: dict[str, Any], datasets: dict[str, Any]) -> None:
        stocks = bundle["stocks"]
        basis = bundle.get("basis_date") or "–"
        ws = wb.create_sheet("한장보고")
        max_col = 20
        self._title_block(
            ws,
            "KOSPI 종목 종가 분석 (한눈에 보기)",
            f"기준일: {basis} (전일종가)  |  기간: 2023년~{str(basis)[:4]}  |  "
            f"{' · '.join(stock['name'] for stock in stocks[:8])}",
            "출처: NAVER Finance / 한국거래소(KRX) 종가·재무  |  참고용 자료 (투자 권유 아님)",
            max_col,
        )
        ws.column_dimensions["A"].width = 22
        for column in range(2, len(stocks) + 2):
            ws.column_dimensions[get_column_letter(column)].width = 16

        self._section_header(ws, 5, "핵심 가격 지표", len(stocks) + 1)
        self._table_header(ws, 6, ["지표"] + [stock["name"] for stock in stocks])
        price_rows = [
            ("최신 종가 (원)", lambda m: m.get("latest"), PRICE_FMT, False),
            ("전일 대비", lambda m: m.get("change_pct"), PERCENT_SIGNED_FMT, True),
            ("2023년말 대비", lambda m: m.get("since_2023_return"), PERCENT_SIGNED_FMT, True),
            ("연초 대비 (YTD)", lambda m: m.get("ytd_return"), PERCENT_SIGNED_FMT, True),
            ("기간 최고가 (원)", lambda m: m.get("high"), PRICE_FMT, False),
            ("기간 최저가 (원)", lambda m: m.get("low"), PRICE_FMT, False),
        ]
        current_row = 7
        for label, getter, number_format, signed in price_rows:
            self._label_cell(ws, current_row, label, fill_color="FFFFFF" if current_row % 2 else ROW_ALT)
            for index, stock in enumerate(stocks, start=2):
                metrics = stock.get("report_metrics") or {}
                value = getter(metrics)
                color = None
                if signed and isinstance(value, (int, float)):
                    color = RISE if value > 0 else FALL if value < 0 else MUTED
                self._value_cell(
                    ws,
                    current_row,
                    index,
                    value if value is not None else "–",
                    number_format=number_format,
                    color=color,
                    bold=signed,
                    fill_color="FFFFFF" if current_row % 2 else ROW_ALT,
                )
            ws.row_dimensions[current_row].height = 21
            current_row += 1

        current_row += 1
        self._section_header(ws, current_row, "밸류에이션·재무 지표", len(stocks) + 1)
        current_row += 1
        self._table_header(ws, current_row, ["지표"] + [stock["name"] for stock in stocks])
        current_row += 1
        fundamental_rows = [
            ("시가총액 (억원)", lambda s: (s.get("market_cap") or 0) / 100_000_000, COUNT_FMT),
            ("PER", lambda s: s.get("per"), RATIO_FMT),
            ("PBR", lambda s: s.get("pbr"), RATIO_FMT),
            ("ROE", lambda s: s.get("roe"), PERCENT_FMT),
            ("영업이익률", lambda s: s.get("op_margin"), PERCENT_FMT),
            ("부채비율", lambda s: s.get("debt_ratio"), PERCENT_FMT),
            ("배당수익률", lambda s: s.get("dividend_yield"), PERCENT_FMT),
            ("매출액 (억원)", lambda s: s.get("sales"), COUNT_FMT),
            ("당기순이익 (억원)", lambda s: s.get("net_income"), COUNT_FMT),
        ]
        for offset, (label, getter, number_format) in enumerate(fundamental_rows):
            fill_color = "FFFFFF" if offset % 2 == 0 else ROW_ALT
            self._label_cell(ws, current_row, label, fill_color=fill_color)
            for index, stock in enumerate(stocks, start=2):
                value = getter(stock)
                self._value_cell(
                    ws,
                    current_row,
                    index,
                    value if value is not None else "–",
                    number_format=number_format,
                    fill_color=fill_color,
                )
            ws.row_dimensions[current_row].height = 21
            current_row += 1
        # Fiscal-year footnote row
        self._label_cell(ws, current_row, "재무 데이터 기준", fill_color=PALE_AMBER)
        for index, stock in enumerate(stocks, start=2):
            year = stock.get("fiscal_year")
            self._value_cell(
                ws,
                current_row,
                index,
                f"{year} 확정" if year else "데이터 없음",
                alignment=CENTER,
                color=NAVY_2,
                fill_color=PALE_AMBER,
                bold=True,
            )
        ws.row_dimensions[current_row].height = 22

        chart_row = current_row + 3
        ws.merge_cells(start_row=chart_row - 1, start_column=1, end_row=chart_row - 1, end_column=max_col)
        heading = ws.cell(chart_row - 1, 1, "그래프 ① 일별 종가 · ② 월별 종가 · ③ 연간 수익률")
        heading.font = font(11, True, NAVY_2)
        heading.fill = fill(PALE_BLUE)
        heading.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        heading.border = BORDER
        ws.row_dimensions[chart_row - 1].height = 24

        daily_chart = self._line_chart(
            wb["일일종가"],
            "① 일별 종가 추이 (전체 기간)",
            stocks,
            header_row=4,
            first_row=5,
            last_row=4 + len(datasets["dates"]),
            y_title="종가 (원)",
            tick_skip=max(1, len(datasets["dates"]) // 8),
        )
        daily_chart.width = 25.5
        daily_chart.height = 9.6
        ws.add_chart(daily_chart, f"A{chart_row}")

        monthly_chart = self._line_chart(
            wb["월별종가"],
            "② 월별 종가 추이",
            stocks,
            header_row=4,
            first_row=5,
            last_row=4 + len(datasets["months"]),
            y_title="종가 (원)",
            tick_skip=max(1, len(datasets["months"]) // 6),
        )
        monthly_chart.width = 12.4
        monthly_chart.height = 8.5
        ws.add_chart(monthly_chart, f"A{chart_row + 20}")

        annual_chart = self._annual_bar_chart(
            wb["연간데이터"],
            "③ 연간 수익률 비교",
            stocks,
            years=datasets["years"],
        )
        annual_chart.width = 12.4
        annual_chart.height = 8.5
        ws.add_chart(annual_chart, f"H{chart_row + 20}")

        note_row = chart_row + 40
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=max_col)
        note = ws.cell(
            note_row,
            1,
            "※ 그래프는 종목별 종가 데이터를 기반으로 생성됩니다. "
            "전일 종가 데이터는 2023-01-02부터 직전 거래일까지이며, "
            "재무 지표는 최근 확정 연간 실적, PER·배당은 네이버증권 최신 값을 사용합니다.  "
            + DISCLAIMER,
        )
        note.font = font(8, False, MUTED)
        note.alignment = WRAP_LEFT
        note.fill = fill(ROW_ALT)
        note.border = BORDER
        ws.row_dimensions[note_row].height = 30
        ws.freeze_panes = "A7"

    # ------------------------------------------------------------------
    # Sheet 2: 전체보고
    # ------------------------------------------------------------------
    def _build_detail_sheet(self, wb: Workbook, bundle: dict[str, Any], datasets: dict[str, Any]) -> None:
        stocks = bundle["stocks"]
        basis = bundle.get("basis_date") or "–"
        ws = wb.create_sheet("전체보고")
        max_col = 20
        self._title_block(
            ws,
            "KOSPI 종목 종가 분석 전체 보고서",
            f"상세 성과·재무 비교  |  기준일: {basis}  |  기간: 2023-01-01 ~ {basis}",
            "출처: NAVER Finance / 한국거래소(KRX) | 단위: 원, %, 억원 | 투자 권유 아님",
            max_col,
        )
        self._section_header(ws, 5, "종목별 상세 성과", max_col)
        headers = [
            "종목명",
            "코드",
            "최신 종가(원)",
            "전일 대비",
            "2023년말 대비",
            "연초 대비(YTD)",
            "기간 최고(원)",
            "기간 최저(원)",
            "52주 최고(원)",
            "52주 최저(원)",
            "시가총액(억원)",
            "PER(배)",
            "PBR(배)",
            "ROE",
            "영업이익률",
            "부채비율",
            "배당수익률",
            "매출액(억원)",
            "당기순이익(억원)",
            "재무 기준",
        ]
        self._table_header(ws, 6, headers)
        for index, stock in enumerate(stocks, start=7):
            metrics = stock.get("report_metrics") or {}
            values = [
                stock.get("name"),
                stock.get("code"),
                metrics.get("latest"),
                metrics.get("change_pct"),
                metrics.get("since_2023_return"),
                metrics.get("ytd_return"),
                metrics.get("high"),
                metrics.get("low"),
                stock.get("high52"),
                stock.get("low52"),
                round((stock.get("market_cap") or 0) / 100_000_000, 1),
                stock.get("per"),
                stock.get("pbr"),
                stock.get("roe"),
                stock.get("op_margin"),
                stock.get("debt_ratio"),
                stock.get("dividend_yield"),
                stock.get("sales"),
                stock.get("net_income"),
                f"{stock.get('fiscal_year')} 확정" if stock.get("fiscal_year") else "–",
            ]
            row_fill = fill("FFFFFF" if (index - 7) % 2 == 0 else ROW_ALT)
            for column, value in enumerate(values, start=1):
                cell = ws.cell(index, column, value if value is not None else "–")
                cell.font = font(9)
                cell.fill = row_fill
                cell.border = BORDER
                cell.alignment = LEFT if column == 1 else CENTER
                if column in (3, 7, 8, 9, 10):
                    cell.number_format = PRICE_FMT
                    cell.alignment = RIGHT
                elif column in (4, 5, 6):
                    cell.number_format = PERCENT_SIGNED_FMT
                    if isinstance(value, (int, float)):
                        cell.font = font(9, True, RISE if value > 0 else FALL if value < 0 else MUTED)
                elif column in (14, 15, 16, 17):
                    cell.number_format = PERCENT_FMT
                elif column in (11, 18, 19):
                    cell.number_format = COUNT_FMT
                elif column in (12, 13):
                    cell.number_format = RATIO_FMT
                if column == 2:
                    cell.hyperlink = Hyperlink(ref=cell.coordinate, target=naver_stock_url(stock["code"]))
                    cell.font = font(9, False, FALL)
            ws.row_dimensions[index].height = 22
        widths = [18, 9, 14, 10, 13, 13, 13, 13, 13, 13, 16, 9, 9, 9, 12, 11, 12, 14, 15, 12]
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.freeze_panes = "C7"

        chart_start = 6 + len(stocks) + 3
        ws.merge_cells(start_row=chart_start - 1, start_column=1, end_row=chart_start - 1, end_column=max_col)
        heading = ws.cell(chart_start - 1, 1, "상세 차트")
        heading.font = font(11, True, NAVY_2)
        heading.fill = fill(PALE_BLUE)
        heading.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        heading.border = BORDER

        chart1 = self._line_chart(
            wb["일일종가"],
            "일별 종가 추이 (전체 기간)",
            stocks,
            header_row=4,
            first_row=5,
            last_row=4 + len(datasets["dates"]),
            y_title="종가 (원)",
            tick_skip=max(1, len(datasets["dates"]) // 9),
        )
        chart1.width = 26
        chart1.height = 10.6
        ws.add_chart(chart1, f"A{chart_start}")

        chart2 = self._line_chart(
            wb["월별종가"],
            "월별 종가 추이",
            stocks,
            header_row=4,
            first_row=5,
            last_row=4 + len(datasets["months"]),
            y_title="종가 (원)",
            tick_skip=max(1, len(datasets["months"]) // 7),
        )
        chart2.width = 12.6
        chart2.height = 9.2
        ws.add_chart(chart2, f"A{chart_start + 23}")

        chart3 = self._annual_bar_chart(
            wb["연간데이터"],
            "연간 수익률 비교",
            stocks,
            years=datasets["years"],
        )
        chart3.width = 12.6
        chart3.height = 9.2
        ws.add_chart(chart3, f"H{chart_start + 23}")

        note_row = chart_start + 45
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=max_col)
        note = ws.cell(
            note_row,
            1,
            "※ 2023년말 대비 수익률은 2023년 마지막 거래일 종가 기준, 연초 대비는 직전 연도 마지막 거래일 종가 기준입니다. "
            "재무 지표는 최근 확정 연간 실적, PER·PBR·배당은 네이버증권 최신 값입니다. " + DISCLAIMER,
        )
        note.font = font(8, False, MUTED)
        note.alignment = WRAP_LEFT
        note.fill = fill(ROW_ALT)
        note.border = BORDER
        ws.row_dimensions[note_row].height = 30

    # ------------------------------------------------------------------
    # Sheet 3: 일일종가
    # ------------------------------------------------------------------
    def _build_daily_sheet(self, wb: Workbook, bundle: dict[str, Any], datasets: dict[str, Any]) -> None:
        stocks = bundle["stocks"]
        dates = datasets["dates"]
        lookup = datasets["lookup"]
        ws = wb.create_sheet("일일종가")
        max_col = len(stocks) + 1
        self._title_block(
            ws,
            f"KOSPI 종목 일일 종가 ({dates[0] if dates else '–'} ~ {dates[-1] if dates else '–'})",
            "출처: NAVER Finance / 한국거래소(KRX) | 단위: 원(KRW) | 전일 종가 기준",
            f"기준일: {bundle.get('basis_date') or '–'}  |  데이터 행수: {len(dates):,}  |  종목수: {len(stocks)}",
            max_col,
        )
        self._table_header(ws, 4, ["날짜"] + [f"{stock['name']} ({stock['code']})" for stock in stocks])
        ws.freeze_panes = "B5"
        ws.column_dimensions["A"].width = 13
        for index in range(2, max_col + 1):
            ws.column_dimensions[get_column_letter(index)].width = 17

        for row_index, day in enumerate(dates, start=5):
            day_cell = ws.cell(row_index, 1, day)
            day_cell.font = font(9, False, TEXT)
            day_cell.alignment = CENTER
            day_cell.border = BORDER
            day_cell.fill = fill("FFFFFF" if row_index % 2 == 0 else ROW_ALT)
            for column_index, stock in enumerate(stocks, start=2):
                value = lookup[stock["code"]].get(day)
                cell = ws.cell(row_index, column_index, value)
                cell.font = font(9, False)
                cell.alignment = RIGHT
                cell.border = BORDER
                cell.fill = fill("FFFFFF" if row_index % 2 == 0 else ROW_ALT)
                cell.number_format = PRICE_FMT
        last_data_row = 4 + len(dates)
        ws.auto_filter.ref = f"A4:{get_column_letter(max_col)}{last_data_row}"
        summary_start = last_data_row + 2
        for offset, (label, func) in enumerate((("기간 최고가", "MAX"), ("기간 최저가", "MIN"))):
            row = summary_start + offset
            label_cell = ws.cell(row, 1, label)
            label_cell.font = font(9, True, NAVY_2)
            label_cell.fill = fill(PALE_AMBER)
            label_cell.alignment = LEFT
            label_cell.border = BORDER
            for column_index in range(2, max_col + 1):
                cell = ws.cell(
                    row,
                    column_index,
                    f"={func}({get_column_letter(column_index)}5:{get_column_letter(column_index)}{last_data_row})",
                )
                cell.font = font(9, True)
                cell.fill = fill(PALE_AMBER)
                cell.border = BORDER
                cell.alignment = RIGHT
                cell.number_format = PRICE_FMT
            ws.row_dimensions[row].height = 21
        # Latest close row (static, avoids assumptions about equal listing dates).
        latest_row = summary_start + 2
        label_cell = ws.cell(latest_row, 1, "전일 종가")
        label_cell.font = font(9, True, "FFFFFF")
        label_cell.fill = fill(TEAL_DARK)
        label_cell.alignment = LEFT
        label_cell.border = BORDER
        for column_index, stock in enumerate(stocks, start=2):
            metrics = stock.get("report_metrics") or {}
            cell = ws.cell(latest_row, column_index, metrics.get("latest"))
            cell.font = font(9, True, "FFFFFF")
            cell.fill = fill(TEAL_DARK)
            cell.border = BORDER
            cell.alignment = RIGHT
            cell.number_format = PRICE_FMT
        ws.row_dimensions[latest_row].height = 22

    # ------------------------------------------------------------------
    # Sheet 4: 월별종가
    # ------------------------------------------------------------------
    def _build_monthly_sheet(self, wb: Workbook, bundle: dict[str, Any], datasets: dict[str, Any]) -> None:
        stocks = bundle["stocks"]
        months = datasets["months"]
        month_lookup = datasets["month_lookup"]
        ws = wb.create_sheet("월별종가")
        max_col = len(stocks) + 1
        self._title_block(
            ws,
            "KOSPI 종목 월별 종가",
            "각 월 마지막 거래일 종가 기준  |  단위: 원(KRW)",
            f"기간: {months[0] if months else '–'} ~ {months[-1] if months else '–'}  |  데이터: 일일종가 시트 집계",
            max_col,
        )
        self._table_header(ws, 4, ["연-월"] + [f"{stock['name']} ({stock['code']})" for stock in stocks])
        ws.freeze_panes = "B5"
        ws.column_dimensions["A"].width = 13
        for index in range(2, max_col + 1):
            ws.column_dimensions[get_column_letter(index)].width = 17
        for row_index, month in enumerate(months, start=5):
            cell = ws.cell(row_index, 1, month)
            cell.font = font(9)
            cell.alignment = CENTER
            cell.border = BORDER
            cell.fill = fill("FFFFFF" if row_index % 2 == 0 else ROW_ALT)
            for column_index, stock in enumerate(stocks, start=2):
                value = month_lookup[stock["code"]].get(month)
                value_cell = ws.cell(row_index, column_index, value)
                value_cell.font = font(9)
                value_cell.alignment = RIGHT
                value_cell.border = BORDER
                value_cell.fill = fill("FFFFFF" if row_index % 2 == 0 else ROW_ALT)
                value_cell.number_format = PRICE_FMT
        ws.auto_filter.ref = f"A4:{get_column_letter(max_col)}{4 + len(months)}"

    # ------------------------------------------------------------------
    # Sheet 5: 연간데이터
    # ------------------------------------------------------------------
    def _build_annual_sheet(self, wb: Workbook, bundle: dict[str, Any], datasets: dict[str, Any]) -> None:
        stocks = bundle["stocks"]
        years = datasets["years"]
        annual_lookup = datasets["annual_lookup"]
        annual_returns = datasets["annual_returns"]
        ws = wb.create_sheet("연간데이터")
        max_col = 1 + len(stocks) * 2
        basis = bundle.get("basis_date")
        self._title_block(
            ws,
            "KOSPI 종목 연간 종가 및 수익률",
            f"연말 종가 기준 수익률  |  {years[-1] if years else '–'}년은 {basis} 기준 YTD",
            "출처: NAVER Finance / 한국거래소(KRX) 종가  |  단위: 원, %",
            max_col,
        )
        headers = ["연도"]
        for stock in stocks:
            headers.extend([f"{stock['name']} 종가(원)", f"{stock['name']} 수익률"])
        self._table_header(ws, 4, headers)
        ws.freeze_panes = "B5"
        ws.column_dimensions["A"].width = 10
        for column in range(2, max_col + 1):
            ws.column_dimensions[get_column_letter(column)].width = 17
        for row_index, year in enumerate(years, start=5):
            cell = ws.cell(row_index, 1, year)
            cell.font = font(9, True, NAVY_2)
            cell.alignment = CENTER
            cell.border = BORDER
            cell.fill = fill(PALE_BLUE)
            for stock_index, stock in enumerate(stocks):
                close_col = 2 + stock_index * 2
                return_col = close_col + 1
                close_value = annual_lookup[stock["code"]].get(year)
                close_cell = ws.cell(row_index, close_col, close_value)
                close_cell.font = font(9)
                close_cell.number_format = PRICE_FMT
                close_cell.alignment = RIGHT
                close_cell.border = BORDER
                close_cell.fill = fill("FFFFFF" if row_index % 2 == 0 else ROW_ALT)
                return_value = annual_returns[stock["code"]].get(year)
                return_cell = ws.cell(
                    row_index,
                    return_col,
                    return_value if return_value is not None else None,
                )
                return_cell.font = font(
                    9,
                    True,
                    RISE if (return_value or 0) > 0 else FALL if (return_value or 0) < 0 else MUTED,
                )
                return_cell.number_format = PERCENT_SIGNED_FMT
                return_cell.alignment = RIGHT
                return_cell.border = BORDER
                return_cell.fill = fill("FFFFFF" if row_index % 2 == 0 else ROW_ALT)
            ws.row_dimensions[row_index].height = 21
        # Simple average annual return summary.
        summary_row = 5 + len(years) + 1
        cell = ws.cell(summary_row, 1, "평균")
        cell.font = font(9, True, "FFFFFF")
        cell.fill = fill(NAVY_2)
        cell.alignment = CENTER
        cell.border = BORDER
        for stock_index in range(len(stocks)):
            close_col = 2 + stock_index * 2
            return_col = close_col + 1
            ws.cell(summary_row, close_col, "").fill = fill(NAVY_2)
            ws.cell(summary_row, close_col).border = BORDER
            avg_cell = ws.cell(
                summary_row,
                return_col,
                f"=AVERAGE({get_column_letter(return_col)}5:{get_column_letter(return_col)}{4 + len(years)})",
            )
            avg_cell.font = font(9, True, "FFFFFF")
            avg_cell.fill = fill(NAVY_2)
            avg_cell.number_format = PERCENT_SIGNED_FMT
            avg_cell.alignment = RIGHT
            avg_cell.border = BORDER

    # ------------------------------------------------------------------
    # Charts
    # ------------------------------------------------------------------
    def _line_chart(
        self,
        data_ws,
        title: str,
        stocks: Sequence[dict[str, Any]],
        *,
        header_row: int,
        first_row: int,
        last_row: int,
        y_title: str,
        tick_skip: int,
    ) -> LineChart:
        chart = LineChart()
        chart.title = title
        chart.style = 10
        chart.y_axis.title = y_title
        chart.x_axis.title = "날짜"
        chart.y_axis.numFmt = "#,##0"
        chart.y_axis.majorGridlines = ChartLines()
        chart.y_axis.majorGridlines.spPr = GraphicalProperties(
            ln=LineProperties(solidFill="E6EDF3", w=6350)
        )
        chart.x_axis.tickLblSkip = max(1, tick_skip)
        chart.x_axis.tickMarkSkip = max(1, tick_skip)
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.height = 8.5
        chart.width = 20
        chart.legend.position = "b"
        chart.legend.overlay = False
        for index, stock in enumerate(stocks):
            reference = Reference(
                data_ws,
                min_col=2 + index,
                min_row=first_row,
                max_row=last_row,
            )
            series = Series(reference, title=stock["name"])
            series.smooth = False
            series.graphicalProperties.line.solidFill = CHART_COLORS[index % len(CHART_COLORS)].lstrip("#")
            series.graphicalProperties.line.width = 22000
            series.graphicalProperties.line.cap = "rnd"
            chart.series.append(series)
        categories = Reference(data_ws, min_col=1, min_row=first_row, max_row=last_row)
        chart.set_categories(categories)
        return chart

    def _annual_bar_chart(
        self,
        data_ws,
        title: str,
        stocks: Sequence[dict[str, Any]],
        *,
        years: Sequence[int],
    ) -> BarChart:
        chart = BarChart()
        chart.type = "col"
        chart.grouping = "clustered"
        chart.overlap = -10
        chart.gapWidth = 70
        chart.title = title
        chart.style = 10
        chart.y_axis.title = "수익률 (%)"
        chart.x_axis.title = "연도"
        chart.y_axis.numFmt = '0.0"%"'
        chart.y_axis.majorGridlines = ChartLines()
        chart.y_axis.majorGridlines.spPr = GraphicalProperties(
            ln=LineProperties(solidFill="E6EDF3", w=6350)
        )
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.height = 8.5
        chart.width = 20
        chart.legend.position = "b"
        chart.legend.overlay = False
        n = len(stocks)
        first_data_row = 6  # first return row after the 2023 base (data starts at row 5)
        last_data_row = 4 + len(years)
        for index, stock in enumerate(stocks):
            col = 3 + index * 2  # return columns: C, E, G, ...
            if last_data_row < first_data_row:
                continue
            reference = Reference(
                data_ws,
                min_col=col,
                min_row=first_data_row,
                max_row=last_data_row,
            )
            series = Series(reference, title=stock["name"])
            series.graphicalProperties.solidFill = CHART_COLORS[index % len(CHART_COLORS)].lstrip("#")
            series.graphicalProperties.line.solidFill = CHART_COLORS[index % len(CHART_COLORS)].lstrip("#")
            chart.series.append(series)
        if last_data_row >= first_data_row:
            categories = Reference(
                data_ws, min_col=1, min_row=first_data_row, max_row=last_data_row
            )
            chart.set_categories(categories)
        chart.dLbls = DataLabelList()
        chart.dLbls.showVal = False
        return chart

    # ------------------------------------------------------------------
    # Screen helpers
    # ------------------------------------------------------------------
    def _screen_subtitle(self, result: ScreenResult) -> str:
        basis = self.service.db.latest_price_date() or "–"
        return (
            f"기준일: {basis}  |  통과 종목: {len(result.rows):,}개  |  "
            f"KOSPI 전체: {result.total:,}개  |  최근 확정 연차 재무 기준"
        )

    def _criteria_text(self, result: ScreenResult) -> str:
        criteria = result.criteria
        if not criteria:
            return "조건 없음"
        parts: list[str] = []
        if criteria.market_cap_min is not None:
            parts.append(f"시가총액 ≥ {criteria.market_cap_min:,.0f}억원")
        if criteria.debt_ratio_max is not None:
            parts.append(f"부채비율 ≤ {criteria.debt_ratio_max:,.1f}%")
        if criteria.op_margin_min is not None:
            parts.append(f"영업이익률 ≥ {criteria.op_margin_min:,.1f}%")
        if criteria.per_min is not None or criteria.per_max is not None:
            low = f"{criteria.per_min:,.1f}" if criteria.per_min is not None else "0"
            high = f"{criteria.per_max:,.1f}" if criteria.per_max is not None else "∞"
            parts.append(f"PER {low}~{high}배")
        if criteria.sales_min is not None:
            parts.append(f"매출액 ≥ {criteria.sales_min:,.0f}억원")
        if criteria.roe_min is not None:
            parts.append(f"ROE ≥ {criteria.roe_min:,.1f}%")
        if criteria.dividend_yield_min is not None:
            parts.append(f"배당수익률 ≥ {criteria.dividend_yield_min:,.1f}%")
        if criteria.net_income_positive:
            parts.append("당기순이익 흑자")
        if not criteria.include_preferred:
            parts.append("우선주 제외")
        if criteria.query:
            parts.append(f"검색어 '{criteria.query}'")
        return "  ·  ".join(parts) if parts else "조건 없음"
