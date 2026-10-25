# -*- coding: utf-8 -*-
"""Developer interaction smoke test (offscreen, temporary database copy)."""

from __future__ import annotations

import os
import shutil
import sys
import tempfile
import time
from datetime import date
from pathlib import Path
from typing import Any, Callable

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
os.environ["KOSPI_INSIGHT_NO_AUTO_REFRESH"] = "1"

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from PySide6.QtCore import QDate, QSettings, Qt  # noqa: E402
from PySide6.QtWidgets import QApplication, QMessageBox  # noqa: E402

from kospi_insight.database import Database  # noqa: E402
from kospi_insight.dialogs import AddStockDialog  # noqa: E402
from kospi_insight.main_window import MainWindow  # noqa: E402
from kospi_insight.sectors import is_financial_stock  # noqa: E402
from kospi_insight.services import FilterCriteria, MarketService  # noqa: E402
from kospi_insight.theme import apply_theme  # noqa: E402

PASSED: list[str] = []
FAILED: list[str] = []


def check(name: str, condition: Any, detail: str = "") -> None:
    if condition:
        PASSED.append(name)
        print(f"  PASS {name}")
    else:
        FAILED.append(f"{name} {detail}".strip())
        print(f"  FAIL {name} {detail}".strip())


def pump(app: QApplication, seconds: float = 0.12) -> None:
    deadline = time.time() + seconds
    while time.time() < deadline:
        app.processEvents()
        time.sleep(0.01)


def main() -> int:
    temp_dir = Path(tempfile.mkdtemp(prefix="kospi_ui_"))
    db_path = temp_dir / "test.db"
    shutil.copy2(ROOT / "data" / "kospi_insight.db", db_path)

    QSettings("KOSPIInsight", "KOSPI Insight").clear()
    app = QApplication(sys.argv[:1])
    app.setApplicationName("KOSPI Insight Test")
    apply_theme(app)

    db = Database(db_path)
    service = MarketService(db)
    service.bootstrap()
    window = MainWindow(service)
    window.resize(1440, 920)
    window.show()
    window.start()
    pump(app, 0.3)

    messages: list[tuple[str, str]] = []
    for page in window.pages.values():
        page.message_requested.connect(lambda text, kind: messages.append((text, kind)))
    window.toast.hide()

    # ------------------------------------------------------------------
    print("\n[1] Screener")
    window.navigate("screener")
    pump(app)
    screener = window.screener
    check("default result rows", len(screener._last_result.rows) == 621, str(len(screener._last_result.rows)))
    check("default criteria cap", screener._last_criteria.market_cap_min == 1000)
    check("default filter badge", screener.filter_count.text() == "필터 1개", screener.filter_count.text())
    check("initial state", screener.table.model().rowCount() == 621)

    for preset in screener.PRESETS:
        screener.apply_preset(preset)
        pump(app, 0.05)
        rows = len(screener._last_result.rows) if screener._last_result else 0
        check(f"preset {preset} has rows", rows > 0, str(rows))

    screener.apply_preset("기본")
    screener._set_share_type("preferred")
    screener.apply_filters()
    pump(app, 0.05)
    preferred_rows = screener._last_result.rows
    check(
        "preferred-only rows",
        bool(preferred_rows) and all(row.get("is_preferred") for row in preferred_rows),
        str(len(preferred_rows)),
    )
    screener._set_share_type("common")
    screener.apply_filters()
    pump(app, 0.05)
    common_rows = screener._last_result.rows
    check(
        "common-only rows",
        bool(common_rows) and not any(row.get("is_preferred") for row in common_rows),
        str(len(common_rows)),
    )
    screener._set_share_type("all")
    screener.financial_check.setChecked(True)
    screener._fields["debt_ratio"].set_text("150")
    screener.apply_filters()
    pump(app, 0.05)
    debt_rows = screener._last_result.rows
    check("debt-ratio rows", bool(debt_rows), str(len(debt_rows)))
    check(
        "financials excluded from debt ratio",
        not any(is_financial_stock(row["code"], row.get("name") or "") for row in debt_rows),
    )
    screener.financial_check.setChecked(False)
    screener.apply_filters()
    pump(app, 0.05)
    check(
        "financials can be included when option off",
        any(
            is_financial_stock(row["code"], row.get("name") or "")
            for row in screener._last_result.rows
        ),
    )
    screener.reset_filters()
    pump(app, 0.05)

    screener.apply_preset("기본")
    screener._fields["per_min"].set_text("20")
    screener._fields["per_max"].set_text("10")
    before = len(screener._last_result.rows)
    messages.clear()
    screener.apply_filters()
    pump(app, 0.05)
    check("invalid PER rejected", len(screener._last_result.rows) == before)
    check("invalid PER message", any("PER" in text for text, _ in messages), str(messages[-2:]))
    screener._fields["per_min"].clear()
    screener._fields["per_max"].clear()
    screener.query_edit.setText("삼성전자")
    screener.apply_filters()
    pump(app, 0.05)
    check(
        "query result includes Samsung",
        any(row["code"] == "005930" for row in screener._last_result.rows),
    )
    screener.reset_filters()
    pump(app, 0.05)
    check("reset rows", len(screener._last_result.rows) == 621, str(len(screener._last_result.rows)))

    # star toggle on a default stock with existing prices
    rows = screener._last_result.rows
    target = next((row["code"] for row in rows if row["code"] == "005930"), None)
    if target is None:
        target = rows[0]["code"]
    in_watch = db.is_in_watchlist(target)
    screener._toggle_star(target)
    check("screener star toggles", db.is_in_watchlist(target) != in_watch)
    if db.is_in_watchlist(target) != in_watch:
        screener._toggle_star(target)
    check("screener star restores", db.is_in_watchlist(target) == in_watch)

    # ------------------------------------------------------------------
    print("\n[2] Analysis")
    window.navigate("analysis", "005930")
    pump(app, 0.25)
    analysis = window.analysis
    check("analysis current", analysis.current_code == "005930")
    check("chart series count", len(analysis.chart.payload.get("series") or []) >= 1)
    check("chart full categories", len(analysis.chart.payload.get("categories") or []) > 100)

    analysis.period_control._buttons["annual"].click()
    pump(app, 0.05)
    check("annual chart mode", analysis.chart.payload.get("period") == "annual")
    check("annual categories years", analysis.chart.payload.get("categories") == ["2023", "2024", "2025", "2026"], str(analysis.chart.payload.get("categories")))
    analysis.period_control._buttons["monthly"].click()
    pump(app, 0.05)
    check("monthly categories", len(analysis.chart.payload.get("categories") or []) >= 30)
    analysis.period_control._buttons["recent"].click()
    pump(app, 0.05)
    check("recent 30 sessions", len(analysis.chart.payload.get("categories") or []) == 30)
    analysis.mode_control._buttons["return"].click()
    pump(app, 0.05)
    check("return mode", analysis.chart.payload.get("mode") == "return")
    check("return unit", analysis.chart.payload.get("unit") == "%")

    analysis.date_edit.setDate(QDate(2026, 10, 23))
    analysis._lookup_current()
    pump(app, 0.05)
    check("lookup latest date", analysis.lookup_date_label.text().startswith("2026-10-23"), analysis.lookup_date_label.text())
    check("lookup price non-empty", analysis.lookup_price.text() not in {"–", ""}, analysis.lookup_price.text())
    analysis.date_edit.setDate(QDate(2026, 10, 3))
    analysis._lookup_current()
    pump(app, 0.05)
    check("holiday falls back", "직전" in analysis.lookup_exact_label.text() or "휴장" in analysis.lookup_exact_label.text(), analysis.lookup_exact_label.text())

    # chart legend toggle does not crash
    if analysis.chart.payload.get("series"):
        code = analysis.chart.payload["series"][0]["code"]
        analysis.chart.hidden_codes = {code}
        analysis.chart.update()
        analysis.chart.hidden_codes = set()
        analysis.chart.update()
        check("legend toggle no crash", True)

    # ------------------------------------------------------------------
    print("\n[3] Favorites")
    window.navigate("favorites")
    pump(app, 0.1)
    favorites = window.favorites
    original = db.watchlist_rows()
    check("favorites rows", favorites.table.rowCount() == len(original), str(favorites.table.rowCount()))
    first = original[0]["code"]
    row_index = next(i for i, row in enumerate(original) if row["code"] == first)
    alias_item = favorites.table.item(row_index, 3)
    memo_item = favorites.table.item(row_index, 4)
    old_alias, old_memo = alias_item.text(), memo_item.text()
    alias_item.setText("테스트별칭")
    memo_item.setText("테스트메모")
    pump(app, 0.05)
    db_row = db.get_watchlist_item(first) if hasattr(db, "get_watchlist_item") else None
    rows_now = {row["code"]: row for row in db.watchlist_rows()}
    check("alias saved", rows_now[first].get("alias") == "테스트별칭", str(rows_now[first].get("alias")))
    check("memo saved", rows_now[first].get("memo") == "테스트메모", str(rows_now[first].get("memo")))

    # restore
    rows_now = {row["code"]: row for row in db.watchlist_rows()}
    favorites.refresh()
    pump(app, 0.05)
    idx = next(i for i, row in enumerate(db.watchlist_rows()) if row["code"] == first)
    favorites.table.item(idx, 3).setText(old_alias)
    favorites.table.item(idx, 4).setText(old_memo)
    pump(app, 0.05)

    order_before = db.watchlist_codes()
    favorites._move(order_before[1], -1)
    pump(app, 0.05)
    check("move up reorders", db.watchlist_codes()[0] == order_before[1], str(db.watchlist_codes()))
    favorites._move(order_before[1], 1)
    pump(app, 0.05)
    check("move down restores", db.watchlist_codes() == order_before, str(db.watchlist_codes()))

    # Remove last item with auto-confirmed message box, then re-add.
    original_question = QMessageBox.question
    QMessageBox.question = staticmethod(lambda *a, **k: QMessageBox.Yes)  # type: ignore[assignment]
    last = db.watchlist_codes()[-1]
    last_name = (db.get_stock(last) or {}).get("name") or last
    favorites._remove(last, last_name)
    pump(app, 0.05)
    check("remove favorite", not db.is_in_watchlist(last))
    check("empty or table state", favorites.stack.currentIndex() == (0 if db.watchlist_codes() else 1))
    service.add_watchlist(last)
    window._after_watchlist_change(last)
    pump(app, 0.05)
    check("re-add favorite", db.is_in_watchlist(last))
    QMessageBox.question = original_question  # type: ignore[assignment]

    # ------------------------------------------------------------------
    print("\n[4] Add stock dialog")
    dialog = AddStockDialog(service, window)
    dialog.search.setText("005930")
    dialog._populate()
    pump(app, 0.05)
    check("dialog finds Samsung", dialog.selected_code == "005930", str(dialog.selected_code))
    check("dialog already disabled", not dialog.add_button.isEnabled())
    dialog.search.setText("000660")
    dialog._populate()
    pump(app, 0.05)
    check("dialog finds SK Hynix", dialog.selected_code == "000660", str(dialog.selected_code))
    dialog.search.setText("999999")
    dialog._populate()
    pump(app, 0.05)
    check("dialog unknown lookup row", dialog.selected_code == "999999" and dialog.needs_lookup, f"{dialog.selected_code} {dialog.needs_lookup}")
    check("dialog lookup enabled", dialog.add_button.isEnabled())
    dialog.alias_edit.setText("별칭")
    dialog.memo_edit.setText("메모")
    dialog.reject()
    dialog.deleteLater()
    pump(app, 0.05)

    # ------------------------------------------------------------------
    print("\n[5] Global search / navigation")
    window.navigate("dashboard")
    pump(app, 0.05)
    window.search.setText("SK하이닉스")
    window._run_global_search()
    pump(app, 0.1)
    check("search navigates analysis", window._current_page == "analysis")
    check("search opens stock", window.analysis.current_code == "000660", str(window.analysis.current_code))
    window.navigate("settings")
    pump(app, 0.05)
    check("settings page", window._current_page == "settings")
    check("settings has scroll", window.settings.scroll.widget() is not None)

    # ------------------------------------------------------------------
    print("\n[6] Screener service edge cases")
    empty = service.screen(FilterCriteria(market_cap_min=10**9))
    check("impossible filter empty", len(empty.rows) == 0)
    check("empty averages", empty.averages.get("per") is None)
    criteria = FilterCriteria(per_min=5, per_max=10, include_preferred=False)
    result = service.screen(criteria)
    check("preferred exclusion", all(not row.get("is_preferred") for row in result.rows))
    check("per bounds", all(5 <= (row.get("per") or -1) <= 10 for row in result.rows if row.get("per") is not None))

    # ------------------------------------------------------------------
    print("\n[7] Export files")
    from kospi_insight.excel_report import ExcelReport

    report = ExcelReport(service)
    report_path = temp_dir / "report.xlsx"
    screen_path = temp_dir / "screen.xlsx"
    report.export_watchlist(report_path, db.watchlist_codes())
    report.export_screen(screen_path, screener._last_result)
    check("report exists", report_path.exists() and report_path.stat().st_size > 20_000)
    check("screen exists", screen_path.exists() and screen_path.stat().st_size > 20_000)
    try:
        from openpyxl import load_workbook

        wb = load_workbook(report_path)
        check("report sheets", wb.sheetnames == ["한장보고", "전체보고", "일일종가", "월별종가", "연간데이터"], str(wb.sheetnames))
        check(
            "report charts",
            len(wb["한장보고"]._charts) == 3 and len(wb["전체보고"]._charts) == 3,
            f"{len(wb['한장보고']._charts)}/{len(wb['전체보고']._charts)}",
        )
        wb.close()
        wb = load_workbook(screen_path)
        check("screen sheet", wb.sheetnames == ["스크리닝 결과"], str(wb.sheetnames))
        wb.close()
    except Exception as exc:
        check("workbook validation", False, repr(exc))

    window.close()
    app.processEvents()
    shutil.rmtree(temp_dir, ignore_errors=True)

    print("\n============================================")
    print(f"PASSED {len(PASSED)}  FAILED {len(FAILED)}")
    if FAILED:
        for item in FAILED:
            print(" -", item)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
